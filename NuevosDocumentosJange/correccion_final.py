# -*- coding: utf-8 -*-
"""Correccion final consolidada: aplica todos los hallazgos de los 3 agentes
de auditoria, buscando por CONTENIDO (no por indice fijo, que ya demostro
ser inestable) para ser robusto a que el documento haya cambiado de tamano.
Resalta en amarillo todo lo que modifica."""

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run
from docx.shared import Pt

DOC_PATH = Path(r"c:\Users\Remoto\Desktop\TesisPY\NuevosDocumentosJange\Plantilla de Tesis - ARAGON - JANGE - Final.docx")
ANEXO_PATH = Path(r"c:\Users\Remoto\Desktop\TesisPY\NuevosDocumentosJange\ARAGON_JANGE_ANEXO.docx")


def highlight_replace_in_paragraph(paragraph, old_fragment, new_fragment):
    full_text = "".join(r.text for r in paragraph.runs)
    idx = full_text.find(old_fragment)
    if idx == -1:
        return False
    before = full_text[:idx]
    after = full_text[idx + len(old_fragment):]
    ref_elem = paragraph.runs[0]._element
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)

    def add_run(text, highlight):
        new_elem = copy.deepcopy(ref_elem)
        paragraph._p.append(new_elem)
        run = Run(new_elem, paragraph)
        run.text = text
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW if highlight else None
        return run

    if before:
        add_run(before, False)
    add_run(new_fragment, True)
    if after:
        add_run(after, False)
    return True


def fix_by_content(doc, old_fragment, new_fragment, etiqueta, resultados):
    for para in doc.paragraphs:
        full_text = "".join(r.text for r in para.runs)
        if old_fragment in full_text:
            ok = highlight_replace_in_paragraph(para, old_fragment, new_fragment)
            resultados.append((etiqueta, ok))
            return ok
    resultados.append((etiqueta, False))
    return False


def set_cell_highlighted(cell, text):
    old_runs = cell.paragraphs[0].runs
    font_name = old_runs[0].font.name if old_runs else None
    font_size = old_runs[0].font.size if old_runs else Pt(8.5)
    cell.text = str(text)
    run = cell.paragraphs[0].runs[0]
    if font_name:
        run.font.name = font_name
    run.font.size = font_size or Pt(8.5)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


# =================================================================
# DOCUMENTO PRINCIPAL
# =================================================================
doc = Document(DOC_PATH)
print(f"Documento principal: {len(doc.paragraphs)} parrafos, {len(doc.tables)} tablas")

resultados = []

fix_by_content(doc,
    "mantiene una capacidad predictiva igual o superior a la del subconjunto de 4 variables confirmadas",
    "mantiene una capacidad predictiva igual o superior a la del subconjunto de 6 variables confirmadas",
    "HE1: 4->6 confirmadas", resultados)

fix_by_content(doc,
    "El resultado de este segundo filtro constituye la muestra analítica final: 29,450 registros (68,262 excluidos adicionales)",
    "El resultado de este segundo filtro constituye la muestra analítica final: 39,110 registros (58,602 excluidos adicionales)",
    "4.6 muestra final 29450->39110", resultados)

fix_by_content(doc,
    "se obtuvo un dataset consolidado de 29,450 registros (108,946 registros eliminados",
    "se obtuvo un dataset consolidado de 39,110 registros (99,286 registros eliminados",
    "4.6 dataset consolidado", resultados)

fix_by_content(doc,
    "80% (23,560 registros) destinado al entrenamiento de los modelos y 20% (5,890 registros)",
    "80% (31,288 registros) destinado al entrenamiento de los modelos y 20% (7,822 registros)",
    "4.6 particion 80/20", resultados)

fix_by_content(doc,
    "proporción similar (8.95% con pérdida / 91.05% sin pérdida)",
    "proporción similar (8.75% con pérdida / 91.25% sin pérdida)",
    "4.6 proporcion clases", resultados)

fix_by_content(doc,
    "De las 20 variables evaluadas Boruta confirmo estadísticamente solo 4 como relevantes, altitud conglomerada (V040), intervalo intergenésico (V222), índice de masa corporal (IMC), índice de riqueza (V190). Las 16 variables restantes",
    "De las 19 variables evaluadas Boruta confirmo estadísticamente 6 como relevantes: paridad (V201), intervalo intergenésico (V222), altitud conglomerada (V040), edad (V012), índice de masa corporal (IMC), índice de riqueza (V190). Las 13 variables restantes",
    "4.10.3 boruta 4->6, 20->19", resultados)

fix_by_content(doc,
    "Tabla 11: Las 10 variables retenidas para el modelado, de un total de 20 evaluadas por Boruta",
    "Tabla 11: Las 10 variables retenidas para el modelado, de un total de 19 evaluadas por Boruta",
    "leyenda Tabla 11", resultados)

fix_by_content(doc,
    "Figura 4: Importancia de las 20 variables candidatas según Boruta",
    "Figura 4: Importancia de las 19 variables candidatas según Boruta",
    "leyenda Figura 4", resultados)

fix_by_content(doc,
    "[edad, altitud, intervalo intergenésico, IMC, índice de riqueza, paridad, controles prenatales, nivel educativo, área de residencia]",
    "[edad, altitud, intervalo intergenésico, IMC, índice de riqueza, paridad, consumo de alcohol (QS206), nivel educativo, área de residencia]",
    "lista variables numericas (M14->QS206)", resultados)

fix_by_content(doc,
    "Los resultados se resumen en la Tabla 5.",
    "Los resultados se resumen en la Tabla 13.",
    "referencia Tabla 5 -> Tabla 13", resultados)

fix_by_content(doc,
    "pero un Recall de la clase 1 de exactamente 0: no detectó ni un solo caso real de pérdida gestacional, simplemente predijo \"sin pérdida\" para prácticamente todos los casos",
    "pero un Recall de la clase 1 de apenas 8% (0.0804): detectó muy pocos casos reales de pérdida gestacional, prediciendo \"sin pérdida\" para la gran mayoría de los casos",
    "4.10.4 recall 0 -> 8%", resultados)

fix_by_content(doc,
    "De los 2,636 casos reales de pérdida gestacional en el conjunto evaluado, el modelo detectó correctamente 1,284 (Recall ≈ 48.7%) y no detectó 1,352 (51.3%)",
    "De los 3,422 casos reales de pérdida gestacional en el conjunto evaluado, el modelo detectó correctamente 1,892 (Recall ≈ 55.3%) y no detectó 1,530 (44.7%)",
    "4.10.6 matriz confusion narrativa", resultados)

fix_by_content(doc,
    "estos 1,352 falsos negativos representan",
    "estos 1,530 falsos negativos representan",
    "4.10.6 FN narrativa", resultados)

# Random Forest vs Gradient Boosting: el codigo real (shap_global.py) explica un
# Random Forest, no Gradient Boosting. Se corrige el texto para reflejar la
# realidad del pipeline (mas rapido y seguro que re-correr XAI con otro modelo
# antes de la entrega de hoy).
fix_by_content(doc,
    "Por eso, sobre el modelo Gradient Boosting se aplicaron dos niveles de explicabilidad",
    "Por eso, sobre el modelo Random Forest se aplicaron dos niveles de explicabilidad",
    "4.10.7 XAI modelo RF no GB (1)", resultados)

fix_by_content(doc,
    "aunque el modelo real (Gradient Boosting) sea una \"caja negra\" mucho más compleja",
    "aunque el modelo real (Random Forest) sea una \"caja negra\" mucho más compleja",
    "4.10.7 XAI modelo RF no GB (2)", resultados)

# Tabla "Consolidado de Registros del Dataset" (Año/Registros/Columnas)
tabla_consolidado_arreglada = False
for t in doc.tables:
    header = [c.text.strip() for c in t.rows[0].cells]
    if header == ['Año', 'Registros', 'Columnas']:
        datos = [
            ('2022', '10,004', '23'),
            ('2023', '10,010', '23'),
            ('2024', '9,693', '23'),
            ('2025', '9,403', '23'),
            ('Total consolidado', '39,110', '24 (incluye "anio")'),
        ]
        for i, fila in enumerate(datos, start=1):
            for j, val in enumerate(fila):
                set_cell_highlighted(t.rows[i].cells[j], val)
        tabla_consolidado_arreglada = True
        break
resultados.append(("Tabla Consolidado Registros (Año/Registros/Columnas)", tabla_consolidado_arreglada))

# Matriz de Consistencia (documento principal): quitar M14, agregar QS206
matriz_consistencia_arreglada = False
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if 'M14' in cell.text and 'QS900' in cell.text and 'V201' in cell.text:
                old_runs = cell.paragraphs
                # reemplazar el texto completo de la celda, resaltado
                texto_actual = cell.text
                nuevo_texto = texto_actual.replace(
                    "QS900/QS901, QS109, QS200-QS205C, M14, V201, V302/V313, V222",
                    "QS900/QS901, QS109, QS200-QS210, QS206, V201, V302/V313, V222"
                )
                if nuevo_texto == texto_actual:
                    # intentar variante sin el patron exacto: solo quitar M14 y agregar QS206
                    if ", M14," in texto_actual:
                        nuevo_texto = texto_actual.replace(", M14,", ", QS206,")
                if nuevo_texto != texto_actual:
                    cell.text = nuevo_texto
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    matriz_consistencia_arreglada = True
resultados.append(("Matriz de Consistencia: M14 -> QS206 (doc principal)", matriz_consistencia_arreglada))

doc.save(DOC_PATH)

print("\n--- RESULTADOS DOCUMENTO PRINCIPAL ---")
for etiqueta, ok in resultados:
    print(f"   [{'OK' if ok else 'NO ENCONTRADO'}] {etiqueta}")
