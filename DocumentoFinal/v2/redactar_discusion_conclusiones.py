# -*- coding: utf-8 -*-
"""Reemplaza la guia-esqueleto de 6. Discusion y 7. Conclusiones por texto
academico completo. Los bullets/guia viejos quedan tachados en rojo: el
texto nuevo se agrega despues, resaltado en amarillo."""

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor

DOC_PATH = Path(r"c:\Users\Remoto\Desktop\TesisPY\DocumentoFinal\v2\Tesis - ARAGON - JANGE - Final.docx")

doc = Document(DOC_PATH)
print(f"Documento: {len(doc.paragraphs)} parrafos, {len(doc.tables)} tablas")


def mark_red_strike(paragraph):
    for r in paragraph.runs:
        r.font.highlight_color = None
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r.font.strike = True


def insert_after(ref_para, style_name, text):
    new_p = copy.deepcopy(ref_para._p)
    for r in list(new_p.findall(qn('w:r'))):
        new_p.remove(r)
    ref_para._p.addnext(new_p)
    p_obj = Paragraph(new_p, ref_para._parent)
    p_obj.style = doc.styles[style_name]
    run = p_obj.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p_obj


def procesar_seccion(heading_texto, nuevo_texto, es_intro=False):
    """Encuentra heading_texto, marca en rojo el parrafo de guia/bullets
    inmediatamente siguientes (hasta el proximo Heading), e inserta el
    nuevo texto en amarillo despues del ultimo parrafo rojo."""
    idx = None
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() == heading_texto:
            idx = i
            break
    if idx is None:
        print(f"   NO ENCONTRADO: {heading_texto!r}")
        return

    cursor = doc.paragraphs[idx]
    i = idx + 1
    ultimo_rojo = cursor
    while i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        if p.style.name.startswith('Heading'):
            break
        if p.text.strip():
            mark_red_strike(p)
            ultimo_rojo = p
        i += 1

    insert_after(ultimo_rojo, "Normal", nuevo_texto)
    print(f"   OK: {heading_texto!r}")


# =================================================================
# 6. DISCUSION
# =================================================================
print("\n--- 6. Discusion ---")

DISCUSION_6_1 = (
    "Los resultados de Boruta obtenidos en esta investigación (6 de 19 variables confirmadas "
    "estadísticamente) contrastan con los reportados por Tesfa et al. [7], quienes aplicaron el "
    "mismo algoritmo sobre encuestas DHS de África Oriental logrando, en conjunto con SMOTE y "
    "Random Forest, exactitudes superiores al 92% y un AUC de 0.93. Esta diferencia sugiere que la "
    "señal predictiva disponible en la ENDES peruana —una encuesta de hogares sin componente "
    "clínico— es más débil o está más distribuida entre variables que la observada en el contexto "
    "de dicho estudio. Asimismo, el hallazgo de que el criterio híbrido de 10 variables no ofreció "
    "una mejora consistente frente al conjunto de 6 variables confirmadas matiza la utilidad de "
    "ampliar el espacio de características más allá de la evidencia puramente estadística: si bien "
    "la incorporación de variables con respaldo teórico (área de residencia, nivel educativo, "
    "estado civil, consumo de alcohol) no perjudicó el desempeño, tampoco lo mejoró de forma "
    "generalizable entre los 5 modelos evaluados. Esto sugiere que, en investigaciones futuras "
    "sobre esta misma base de datos, podría priorizarse un criterio de selección más parsimonioso "
    "sin sacrificar capacidad predictiva."
)

DISCUSION_6_2 = (
    "El efecto de SMOTE-NC documentado en esta investigación —mejora sostenida en Recall y "
    "F1-Score en los 5 modelos, pero deterioro del PR-AUC y un comportamiento del AUROC que "
    "depende de la familia de algoritmo— no se reporta explícitamente en estudios como el de "
    "Kubahoniyesu y Kabano [4] o Setegn y Dejene [14], donde el balanceo sintético se presenta "
    "como una mejora sin matices sobre la capacidad predictiva general. Este hallazgo constituye un "
    "aporte diferenciador de la presente tesis: evidencia que el sobremuestreo sintético, si bien "
    "corrige eficazmente el sesgo hacia la clase mayoritaria y mejora la sensibilidad clínica del "
    "modelo (una prioridad justificada en el contexto de salud materna), no necesariamente mejora "
    "—e incluso puede reducir levemente— la capacidad de discriminación global del modelo medida "
    "de forma independiente al umbral de decisión."
)

DISCUSION_6_3 = (
    "El AUROC obtenido para el modelo óptimo (Gradient Boosting, AUROC=0.6499) es considerablemente "
    "más modesto que los valores reportados en los estudios internacionales revisados, que alcanzan "
    "AUCs de entre 0.84 y 0.93 [1, 4, 7, 14]. Esta diferencia es coherente con la naturaleza de los "
    "datos empleados: a diferencia de los registros clínicos u hospitalarios utilizados en algunos "
    "de esos estudios, la ENDES es una encuesta de hogares que no incorpora resultados de "
    "laboratorio, historia clínica detallada de gestaciones previas ni factores genéticos, todos "
    "ellos determinantes documentados de la pérdida gestacional. En cuanto a HE3, si bien los "
    "modelos de ensamble (Gradient Boosting, XGBoost, Random Forest) superaron numéricamente a los "
    "modelos tradicionales (SVM, KNN) tanto en AUROC como en F1-Score, las diferencias absolutas "
    "entre el mejor y el peor modelo son de apenas dos a tres centésimas en F1-Score, por lo que "
    "dicha superioridad debe interpretarse como una tendencia consistente más que como una "
    "diferencia de gran magnitud práctica."
)

DISCUSION_6_4 = (
    "El predominio de las variables sociodemográficas (71.5% de la importancia total) sobre las de "
    "salud (28.5%) según TreeSHAP es consistente con lo reportado por Ba et al. [10] y "
    "Meléndez-Asipali et al. [12], quienes también identifican el nivel educativo, el estado civil "
    "y el área de residencia como determinantes relevantes de resultados reproductivos adversos en "
    "poblaciones de bajos y medianos ingresos, incluyendo el contexto peruano. Es relevante señalar "
    "que este ranking, obtenido mediante valores de Shapley sobre el modelo ya entrenado, no "
    "coincide completamente con el obtenido mediante Boruta durante la fase de selección de "
    "características (basado en la disminución media de impureza de un Random Forest de control). "
    "Esta divergencia es un resultado esperado y documentado en la literatura: Strobl et al. [44] "
    "demuestran que la importancia basada en disminución de impureza tiende a sesgarse a favor de "
    "variables con mayor varianza o número de categorías, lo cual puede explicar por qué variables "
    "como el área de residencia o el estado civil —relevantes para SHAP— no fueron confirmadas por "
    "Boruta."
)

DISCUSION_6_5 = (
    "La coincidencia en la dirección del efecto entre LIME y SHAP Waterfall para los tres casos "
    "analizados (V025, V201 y V501) refuerza la confiabilidad interna de las explicaciones "
    "producidas por el modelo, en línea con lo señalado por Setegn y Dejene [14] respecto al valor "
    "de acoplar múltiples técnicas de explicabilidad para validar cruzadamente los hallazgos. El "
    "caso límite —un falso positivo con probabilidad predicha de exactamente 0.5000, en el que "
    "ninguna variable individual mostró una contribución dominante— resulta particularmente valioso "
    "desde una perspectiva práctica: ejemplifica el tipo de perfil ambiguo en el que el modelo tiene "
    "mayor probabilidad de error, información que sería imposible de obtener observando únicamente "
    "la predicción binaria final. No obstante, el análisis se limita a tres casos seleccionados por "
    "su probabilidad predicha, por lo que la consistencia observada no puede generalizarse "
    "estadísticamente sin un análisis local sobre una muestra más amplia."
)

DISCUSION_6_6 = (
    "Más allá de los hallazgos específicos por objetivo, esta investigación presenta limitaciones "
    "metodológicas que deben declararse explícitamente. En primer lugar, el diseño transversal "
    "implica la integración de cuatro rondas independientes de la ENDES (2022-2025) mediante "
    "pooling de datos, sin seguimiento longitudinal de las mismas mujeres. En segundo lugar, dado "
    "que la ENDES no permite distinguir una pérdida gestacional espontánea de una inducida [48], el "
    "desenlace se redefinió como \"pérdida gestacional\" sin el calificativo \"involuntaria\", "
    "declarando esta limitación de forma explícita en lugar de asumir dicha distinción sin "
    "evidencia. En tercer lugar, se identificó un desfase temporal considerable entre el evento de "
    "pérdida y el momento de la entrevista (mediana de aproximadamente 35 meses para los casos de "
    "pérdida temprana), lo que introduce incertidumbre sobre si algunas variables predictoras "
    "reflejan el estado de la mujer en el momento del evento o su situación actual. En cuarto lugar, "
    "la muestra analítica no incorpora los pesos de expansión ni el diseño muestral complejo de la "
    "ENDES, por lo que los hallazgos no deben interpretarse automáticamente como representativos a "
    "nivel nacional. Finalmente, la variable M14 (número de controles prenatales) fue excluida del "
    "modelamiento tras verificarse que, por el diseño del módulo correspondiente de la ENDES, no "
    "podía vincularse de forma confiable al embarazo que terminó en pérdida, lo que constituye una "
    "limitación de disponibilidad de datos que probablemente afecta a otras variables de estado de "
    "salud similares no evaluadas en esta investigación."
)

procesar_seccion("6.1 Sobre la selección de características (OE1)", DISCUSION_6_1)
procesar_seccion("6.2 Sobre el balanceo de clases (OE2)", DISCUSION_6_2)
procesar_seccion("6.3 Sobre el rendimiento comparativo de los modelos (OE3)", DISCUSION_6_3)
procesar_seccion("6.4 Sobre la explicabilidad global (OE4)", DISCUSION_6_4)
procesar_seccion("6.5 Sobre la explicabilidad local (OE5)", DISCUSION_6_5)
procesar_seccion("6.6 Limitaciones metodológicas transversales", DISCUSION_6_6)

# marcar en rojo el parrafo introductorio "[GUIA A DESARROLLAR...]" de 6. Discusion
for i, para in enumerate(doc.paragraphs):
    if para.text.strip().startswith("[GUÍA A DESARROLLAR — este capítulo debe contrastar"):
        mark_red_strike(para)
        break

# =================================================================
# 7. CONCLUSIONES
# =================================================================
print("\n--- 7. Conclusiones ---")

CONCLUSION_7_1 = (
    "El modelamiento predictivo y explicable mediante Machine Learning permitió estimar el riesgo "
    "de pérdida gestacional en mujeres peruanas de 15 a 49 años a partir de factores "
    "sociodemográficos y de salud registrados en la ENDES 2022-2025, con una capacidad de "
    "discriminación estadísticamente significativa pero moderada: el modelo óptimo (Gradient "
    "Boosting) alcanzó un AUROC de 0.6499, cuyo intervalo de confianza al 95% (obtenido mediante "
    "bootstrap no paramétrico) excluye el valor nulo de 0.5. Este resultado confirma que existe una "
    "relación aprovechable entre los factores sociodemográficos y de salud disponibles en la ENDES "
    "y la ocurrencia de pérdida gestacional, aunque su magnitud es modesta en comparación con "
    "estudios que emplean datos clínicos, y no habilita, por sí sola, el uso del modelo como "
    "herramienta de decisión clínica o de política pública sin validación adicional."
)

CONCLUSION_7_2 = (
    "Se implementó y evaluó el algoritmo Boruta para la selección de características, confirmando "
    "estadísticamente 6 de las 19 variables candidatas (paridad, intervalo intergenésico, altitud "
    "del conglomerado, edad, índice de masa corporal e índice de riqueza). El criterio híbrido de "
    "10 variables adoptado para el modelamiento, que incorpora además área de residencia, nivel "
    "educativo, estado civil y consumo de alcohol, no demostró una ventaja predictiva consistente "
    "frente al subconjunto de 6 variables confirmadas, por lo que HE1 se sostiene solo "
    "parcialmente: la reducción de dimensionalidad fue exitosa, pero la ampliación por respaldo "
    "teórico no aportó una mejora generalizable."
)

CONCLUSION_7_3 = (
    "Se evaluó el impacto del balanceo sintético mediante SMOTE-NC, confirmando que incrementa la "
    "sensibilidad (Recall) y el F1-Score en los cinco algoritmos evaluados sin excepción, lo que "
    "sostiene HE2 en su componente central. Sin embargo, se concluye también que este beneficio "
    "tiene una contrapartida: el PR-AUC disminuye de forma consistente en los cinco modelos, y el "
    "AUROC se ve afectado de manera distinta según la familia de algoritmo, por lo que el balanceo "
    "de clases debe entenderse como una herramienta que optimiza la sensibilidad clínica del modelo "
    "en su umbral de decisión, no como una mejora incondicional de su capacidad de discriminación."
)

CONCLUSION_7_4 = (
    "Se comparó el rendimiento predictivo de cinco algoritmos de Machine Learning, concluyendo que "
    "Gradient Boosting constituye el modelo óptimo para este problema (F1=0.2114, AUROC=0.6499, "
    "PR-AUC=0.2050 bajo validación cruzada de 5 particiones). Los algoritmos de ensamble superaron "
    "numéricamente a los modelos tradicionales, sosteniendo HE3, aunque las diferencias observadas "
    "son de magnitud reducida y no deben interpretarse como una superioridad determinante de un "
    "enfoque sobre otro para este problema específico."
)

CONCLUSION_7_5 = (
    "Se determinó la importancia global y la dirección del impacto de las variables predictoras "
    "mediante TreeSHAP, concluyendo que el área de residencia, el estado civil y el nivel educativo "
    "son los determinantes de mayor peso en las predicciones del modelo, y que la dimensión "
    "sociodemográfica concentra una proporción sustancialmente mayor de la importancia total "
    "(71.5%) que la dimensión de salud (28.5%), lo que sostiene la hipótesis específica "
    "correspondiente."
)

CONCLUSION_7_6 = (
    "Se analizó la atribución local del riesgo gestacional en tres perfiles individuales mediante "
    "LIME y SHAP Waterfall, concluyendo que ambas técnicas son consistentes entre sí en la "
    "dirección del efecto de las variables más influyentes, lo que sostiene HE5. El análisis del "
    "caso límite evidenció, además, que la explicabilidad local aporta valor práctico adicional al "
    "identificar perfiles donde el modelo tiene mayor probabilidad de error, información no "
    "disponible a partir de la sola predicción binaria."
)

CONCLUSION_7_7 = (
    "Se recomienda, para investigaciones futuras: (1) explorar módulos de calendario reproductivo "
    "de la ENDES que permitan vincular cada variable predictora al momento específico del embarazo "
    "evaluado, reduciendo el desfase temporal identificado en esta investigación; (2) incorporar el "
    "diseño muestral complejo y los pesos de expansión de la ENDES si se busca realizar inferencia "
    "poblacional a partir de los hallazgos; y (3) evaluar el uso del subconjunto de 6 variables "
    "confirmadas por Boruta como conjunto principal de modelamiento, dado que en esta investigación "
    "igualó el desempeño del criterio híbrido de 10 variables con una menor dimensionalidad."
)

procesar_seccion("7.1 Conclusión general", CONCLUSION_7_1)
procesar_seccion("7.2 Conclusión OE1 (Boruta)", CONCLUSION_7_2)
procesar_seccion("7.3 Conclusión OE2 (SMOTE-NC)", CONCLUSION_7_3)
procesar_seccion("7.4 Conclusión OE3 (Benchmark)", CONCLUSION_7_4)
procesar_seccion("7.5 Conclusión OE4 (SHAP global)", CONCLUSION_7_5)
procesar_seccion("7.6 Conclusión OE5 (LIME/SHAP local)", CONCLUSION_7_6)
procesar_seccion("7.7 Recomendaciones para investigaciones futuras (opcional, separado de las conclusiones)", CONCLUSION_7_7)

for i, para in enumerate(doc.paragraphs):
    if para.text.strip().startswith("[GUÍA A DESARROLLAR — cada conclusión debe responder"):
        mark_red_strike(para)
        break

doc.save(DOC_PATH)
print("\nDiscusion y Conclusiones redactadas completas.")
